"""Synthetic tests for private master connection and application production.

No real personal profile, contact detail, or repo token enters the fixtures.
"""
import json
import stat
import subprocess
import tempfile
import unittest
from pathlib import Path

from resource_miner.application_portfolio import (
    ORDER, assemble_application, build_application,
)
from resource_miner.private_profile_bridge import (
    load_private_master, read_local_portfolio,
)

YAML = """version: '1.0'
profile_id: synthetic-private-profile
contact:
  legal_name: Example Owner
  location: Example City
  email: secret@example.invalid
  phone: '555-0100'
summary:
  professional_brief: 'Program leader building validated project work.'
competencies:
  operations: ['Project delivery', 'Quality systems']
  ai_research: ['AI evaluation']
  data_quality: ['Data pipeline validation']
  technical_skills:
    quality_compliance: ['ISO audits']
    data_analytics: ['Statistical analysis']
    infrastructure: ['GitHub Actions']
achievements:
  - category: Project
    items:
      - title: AI evaluation program
        description: Ran calibration work using documented controls.
      - title: Quality migration program
        description: Tested source integrity during a platform migration.
experience:
  - role: Quality Director
    organization: Sample Hospital
    start_date: '2021-01-01'
    end_date: null
    summary: Directed compliance and quality review.
    bullets: ['Led audit program.']
education:
  - degree: Bachelor of Science
    school: Sample University
publications: []
metadata: {formats: {resume: {audience: 'Job applications'}}}
"""


def _git(path, *args):
    subprocess.run(["git", "-C", str(path), *args],
                   check=True, capture_output=True, text=True)


def make_repo(path: Path, origin: str, files: dict[str, str]):
    path.mkdir()
    _git(path, "init", "-q")
    _git(path, "remote", "add", "origin", origin)
    for name, contents in files.items():
        target = path / name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(contents, encoding="utf-8")
    _git(path, "add", ".")
    _git(path, "-c", "user.name=Test", "-c",
         "user.email=not-a-person@example.invalid",
         "commit", "-qm", "Synthetic fixtures")


class PrivateProfileBridgeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.private = root / "humanaios-internal"
        self.ops = root / "operations"
        self.consumer = root / "opportunity-aggregator"
        self.output = root / "review.docx"
        make_repo(self.private, "git@github.com:humanaios-ui/humanaios-internal.git",
                  {"src/data/master_profile.yaml": YAML})
        make_repo(self.ops, "https://github.com/humanaios-ui/operations.git", {
            "SEED.md": "HumanAIOS LLC. AI behavioral observability research. TRL 2–3.",
        })
        make_repo(self.consumer, "https://github.com/humanaios-ui/opportunity-aggregator", {
            "lib/matching/humanaios-profile.ts": (
                "export const humanaiosProfile = {\n"
                " description: 'Open AI behavioral observability research.',\n"
                " fundingStage: 'research', budgetRange: null,\n"
                " teamProfile: { size: 'solo' },\n"
                "}\n"
                "/**\n * Calculate grant relevance\n */\n"
            )
        })

    def test_private_connection_redacts_identity(self):
        report = read_local_portfolio(self.private, self.ops, self.consumer,
                                      owner_approved=True)
        out = json.dumps(report)
        self.assertNotIn("secret@example.invalid", out)
        self.assertNotIn("Example Owner", out)
        self.assertNotIn("555-0100", out)
        self.assertEqual(report["review_state"], "NO_DRIFT_OBSERVED")
        self.assertEqual(report["authority_effect"], "NONE")
        self.assertEqual(report["read_mode"], "LOCAL_GIT_HEAD_ONLY")
        self.assertFalse(report["eligibility_assessed"])

    def test_requires_explicit_approval(self):
        with self.assertRaises(PermissionError):
            read_local_portfolio(self.private, self.ops, self.consumer)

    def test_rejects_unapproved_repository_origin(self):
        _git(self.private, "remote", "set-url", "origin", "https://github.com/other/private")
        with self.assertRaises(ValueError):
            load_private_master(self.private)

    def test_selected_application_facts_from_master_only(self):
        master, _ = load_private_master(self.private)
        app = assemble_application(master, target="quality migration", max_projects=1)
        self.assertEqual(app["sections"], ORDER)
        self.assertEqual(len(app["projects"]), 1)
        self.assertEqual(app["projects"][0]["heading"], "Quality migration program")
        self.assertIn("achievements[0]", app["projects"][0]["source"])
        self.assertEqual(app["status"], "DRAFT_REVIEW_REQUIRED")

    def test_curated_project_evidence_requires_human_review(self):
        master, _ = load_private_master(self.private)
        master["project_evidence"] = [{
            "title": "Extended clinical research project",
            "mandate": "Build controlled laboratory workflows.",
            "evidence_bullets": ["Authored SOPs and validation plans."],
            "evidence_ref": "owner-supplied-local-docx:sha256-synthetic",
            "review_status": "USER_SUPPLIED_NEEDS_HUMAN_REVIEW",
        }]
        draft = assemble_application(master, target="clinical research", max_projects=10)
        self.assertFalse(any("Extended clinical" in p["heading"] for p in draft["projects"]))
        master["project_evidence"][0]["review_status"] = "APPROVED"
        approved = assemble_application(master, target="clinical research", max_projects=10)
        self.assertTrue(any("Extended clinical" in p["heading"] for p in approved["projects"]))

    def test_owner_private_docx_does_not_overwrite(self):
        from docx import Document
        template = Path(self.temp.name) / "style.docx"
        doc = Document()
        for name in ORDER:
            doc.add_paragraph(name)
        doc.add_table(rows=2, cols=6)
        doc.save(template)
        receipt = build_application(self.private, template, self.output, target="quality",
                                    forbidden_checkouts=(self.ops,))
        self.assertEqual(receipt["status"], "PRIVATE_DRAFT_CREATED")
        self.assertTrue(self.output.exists())
        self.assertEqual(stat.S_IMODE(self.output.stat().st_mode), 0o600)
        text = "\n".join(p.text for p in Document(self.output).paragraphs)
        self.assertIn("SELECTED PROJECT EVIDENCE PORTFOLIO", text)
        self.assertIn("Sample Hospital", text)
        with self.assertRaises(FileExistsError):
            build_application(self.private, template, self.output)

    def test_rejects_private_output_inside_checkout(self):
        from docx import Document
        template = Path(self.temp.name) / "style.docx"
        Document().save(template)
        with self.assertRaises(ValueError):
            build_application(self.private, template, self.private / "leak.docx")


if __name__ == "__main__":
    unittest.main()
