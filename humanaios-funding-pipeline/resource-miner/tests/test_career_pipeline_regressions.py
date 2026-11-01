"""Regression tests for the integrated career pipeline (#997 + #1029).

Synthetic fixtures only. Covers missing sources, dependent-claim downgrading,
and the human-admission boundary: nothing here prepares, approves or sends a
real application.
"""
import json
import os
import shutil
import stat
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # sibling test helpers, any run mode

from resource_miner.application_portfolio import ORDER, build_application
from resource_miner.job_application import downgrade_for_removed_evidence, prepare
from resource_miner.private_profile_bridge import (
    _git_file, load_private_master, project_consumer, read_local_portfolio,
)
from test_private_profile_bridge import YAML, _git, make_repo

ROOT = Path(__file__).resolve().parents[1]


def record(**over):
    base = {
        "application_id": "a1", "employer": "Example", "role": "Role",
        "job_text": "text", "job_source": "https://example.invalid/job",
        "evidence": [
            {"id": "e1", "claim": "Led audits", "source": "doc-1", "state": "DOCUMENTED"},
            {"id": "e2", "claim": "Ran program", "source": "doc-2", "state": "USER_ATTESTED"},
        ],
        "requirements": [
            {"id": "r1", "requirement": "Audit", "kind": "duty", "status": "SUPPORTED",
             "mapping": "DIRECT", "evidence_ids": ["e1", "e2"]},
            {"id": "r2", "requirement": "Programs", "kind": "minimum", "status": "SUPPORTED",
             "mapping": "DIRECT", "evidence_ids": ["e2"]},
            {"id": "r3", "requirement": "Unrelated", "kind": "preferred", "status": "SUPPORTED",
             "mapping": "TRANSFERABLE", "evidence_ids": ["e1"]},
        ],
    }
    base.update(over)
    return base


class MissingSourceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.private, self.ops = root / "humanaios-internal", root / "operations"
        self.consumer = root / "opportunity-aggregator"
        make_repo(self.private, "git@github.com:humanaios-ui/humanaios-internal.git",
                  {"src/data/master_profile.yaml": YAML})
        make_repo(self.ops, "https://github.com/humanaios-ui/operations.git",
                  {"SEED.md": "HumanAIOS LLC. AI behavioral observability research. TRL 2–3."})
        self.ts = ("export const humanaiosProfile = {\n description: %s,\n"
                   " fundingStage: 'research', budgetRange: null,\n teamProfile: { size: 'solo' },\n}\n")
        self.root = root

    def _consumer(self, description="'Open research.'"):
        make_repo(self.consumer, "https://github.com/humanaios-ui/opportunity-aggregator",
                  {"lib/matching/humanaios-profile.ts": self.ts % description})

    def test_missing_consumer_repository_is_a_clean_value_error(self):
        with self.assertRaises(ValueError) as ctx:
            read_local_portfolio(self.private, self.ops, self.root / "absent", owner_approved=True)
        self.assertEqual(str(ctx.exception), "repository checkout not found")

    def test_cli_missing_consumer_has_no_traceback(self):
        result = subprocess.run(
            [sys.executable, "-m", "resource_miner.cli", "connect-private-profile",
             "--private-checkout", str(self.private), "--operations-checkout", str(self.ops),
             "--consumer-checkout", str(self.root / "absent"), "--owner-approved"],
            cwd=ROOT, capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertNotIn("Traceback", result.stderr)
        self.assertIn("repository checkout not found", result.stderr)

    def test_missing_master_file_fails_closed(self):
        _git(self.private, "rm", "-q", "src/data/master_profile.yaml")
        _git(self.private, "-c", "user.name=T", "-c", "user.email=t@example.invalid", "commit", "-qm", "rm")
        with self.assertRaises(ValueError):
            load_private_master(self.private)

    def test_approval_is_checked_before_any_private_read(self):
        # A missing private checkout would raise ValueError if it were read first.
        with self.assertRaises(PermissionError):
            read_local_portfolio(self.root / "absent", self.ops, self.consumer)

    def test_malformed_yaml_error_does_not_leak_source(self):
        bad = YAML.replace("secret@example.invalid", "secret@example.invalid: [unclosed")
        (self.private / "src/data/master_profile.yaml").write_text(bad, encoding="utf-8")
        _git(self.private, "add", ".")
        _git(self.private, "-c", "user.name=T", "-c", "user.email=t@example.invalid", "commit", "-qm", "bad")
        with self.assertRaises(ValueError) as ctx:
            load_private_master(self.private)
        self.assertNotIn("secret@example.invalid", str(ctx.exception))
        self.assertIsNone(ctx.exception.__cause__)
        self.assertTrue(ctx.exception.__suppress_context__)

    def test_escaped_quote_description_is_rejected_not_truncated(self):
        self._consumer(r"'Don\'t claim 374K verified workers'")
        with self.assertRaises(ValueError):
            project_consumer(self.consumer)

    def test_revision_is_the_blob_that_was_read(self):
        revision, data = _git_file(self.private, "humanaios-ui/humanaios-internal",
                                   "src/data/master_profile.yaml")
        blob = subprocess.run(["git", "-C", str(self.private), "cat-file", "blob", revision],
                              capture_output=True, text=True, check=True).stdout.strip()
        self.assertEqual(data, blob)


class OutputBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.private = self.root / "humanaios-internal"
        make_repo(self.private, "git@github.com:humanaios-ui/humanaios-internal.git",
                  {"src/data/master_profile.yaml": YAML})
        from docx import Document
        self.template = self.root / "style.docx"
        doc = Document()
        for name in ORDER:
            doc.add_paragraph(name)
        doc.add_table(rows=2, cols=6)
        doc.save(self.template)

    def test_refuses_any_git_work_tree_even_when_unlisted(self):
        other = self.root / "unlisted-public-checkout"
        make_repo(other, "https://github.com/humanaios-ui/opportunity-aggregator", {"README.md": "x"})
        with self.assertRaises(ValueError):
            build_application(self.private, self.template, other / "leak.docx")
        self.assertFalse((other / "leak.docx").exists())

    def test_output_outside_git_is_created_with_mode_0600(self):
        out_dir = self.root / "outside"
        out_dir.mkdir()
        # TemporaryDirectory may sit inside no repo; skip rather than assert if it does.
        probe = subprocess.run(["git", "-C", str(out_dir), "rev-parse", "--is-inside-work-tree"],
                               capture_output=True, text=True)
        if probe.returncode == 0:
            self.skipTest("temp dir is inside a git work tree on this machine")
        out = out_dir / "app.docx"
        receipt = build_application(self.private, self.template, out)
        self.assertEqual(stat.S_IMODE(out.stat().st_mode), 0o600)
        self.assertEqual(receipt["authority_effect"], "NONE")
        self.assertEqual(receipt["status"], "PRIVATE_DRAFT_CREATED")
        with self.assertRaises(FileExistsError):
            build_application(self.private, self.template, out)

    def test_cli_render_requires_owner_approval(self):
        result = subprocess.run(
            [sys.executable, "-m", "resource_miner.cli", "render-private-application",
             "--private-checkout", str(self.private), "--template", str(self.template),
             "--output", str(self.root / "x.docx")],
            cwd=ROOT, capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("owner approval required", result.stderr)
        self.assertFalse((self.root / "x.docx").exists())


class DependentClaimTests(unittest.TestCase):
    def test_removing_one_source_downgrades_only_dependent_claims(self):
        out = downgrade_for_removed_evidence(record(), ["e2"])
        rows = {r["id"]: r for r in out["requirements"]}
        self.assertEqual(rows["r1"]["status"], "PARTIAL")      # keeps e1
        self.assertEqual(rows["r1"]["evidence_ids"], ["e1"])
        self.assertEqual((rows["r2"]["status"], rows["r2"]["mapping"]), ("UNKNOWN", "GAP"))
        self.assertEqual(rows["r2"]["evidence_ids"], [])
        self.assertEqual(rows["r3"]["status"], "SUPPORTED")    # never cited e2
        self.assertNotIn("e2", [e["id"] for e in out["evidence"]])
        prepare(out)  # the downgraded record is still valid, nothing dangling

    def test_downgrade_does_not_mutate_input(self):
        original = record()
        snapshot = json.dumps(original, sort_keys=True)
        downgrade_for_removed_evidence(original, ["e1", "e2"])
        self.assertEqual(json.dumps(original, sort_keys=True), snapshot)

    def test_in_progress_evidence_cannot_fully_support(self):
        rec = record(evidence=[{"id": "e1", "claim": "c", "source": "s", "state": "IN_PROGRESS"}],
                     requirements=[{"id": "r1", "requirement": "x", "kind": "duty",
                                    "status": "SUPPORTED", "mapping": "DIRECT", "evidence_ids": ["e1"]}])
        with self.assertRaises(ValueError):
            prepare(rec)
        rec["requirements"][0]["status"] = "PARTIAL"
        self.assertEqual(prepare(rec)["application"]["requirements"][0]["status"], "PARTIAL")

    def test_missing_or_contradicted_evidence_stays_distinct(self):
        rec = record(requirements=[
            {"id": "r1", "requirement": "x", "kind": "duty", "status": "CONTRADICTED",
             "mapping": "GAP", "gap": "conflicting source"},
            {"id": "r2", "requirement": "y", "kind": "duty", "status": "UNKNOWN", "mapping": "GAP"},
        ])
        statuses = [r["status"] for r in prepare(rec)["application"]["requirements"]]
        self.assertEqual(statuses, ["CONTRADICTED", "UNKNOWN"])


class CatalogHoldTests(unittest.TestCase):
    def test_conflicting_holds_become_recheck_required(self):
        url = "https://example.invalid/job"
        rows = [{"url": url, "source_verification_status": "LINK_FAILED_404"},
                {"url": url, "source_verification_status": "ROLE_NOT_LISTED"}]
        out = prepare(record(), rows)["application"]
        self.assertEqual(out["source_verification_status"], "RECHECK_REQUIRED")
        out2 = prepare(record(), list(reversed(rows)))["application"]
        self.assertEqual(out2["source_verification_status"], "RECHECK_REQUIRED")

    def test_single_hold_still_overrides_user_import(self):
        url = "https://example.invalid/job"
        rows = [{"url": url, "source_verification_status": "LINK_FAILED_404"}]
        out = prepare(record(source_verification_status="UNVERIFIED"), rows)["application"]
        self.assertEqual(out["source_verification_status"], "LINK_FAILED_404")
        self.assertTrue(out["source_hold"])


class HumanAdmissionBoundaryTests(unittest.TestCase):
    def test_caller_cannot_self_approve_or_mark_submitted(self):
        rec = record(external_state="SUBMITTED", authority="Z2", review_state="APPROVED",
                     availability="VERIFIED", evidence_verification="INDEPENDENTLY_VERIFIED")
        result = prepare(rec)
        app = result["application"]
        self.assertEqual(app["external_state"], "NOT_SUBMITTED")
        self.assertEqual(app["authority"], "NONE")
        self.assertEqual(app["review_state"], "DRAFT_REQUIRES_HUMAN_REVIEW")
        self.assertEqual(app["availability"], "UNVERIFIED")
        self.assertEqual(app["evidence_verification"], "USER_SUPPLIED_NOT_INDEPENDENTLY_CHECKED")
        self.assertIn("Do not submit or contact anyone", result["skill_handoff"])

    def test_caller_cannot_claim_live_verification(self):
        with self.assertRaises(ValueError):
            prepare(record(source_verification_status="VERIFIED_LIVE"))

    def test_missing_job_specification_is_rejected_not_tailored(self):
        with self.assertRaises(ValueError):
            prepare(record(job_text=""))
        with self.assertRaises(ValueError):
            prepare(record(requirements=[]))

    def test_job_application_module_makes_no_network_import(self):
        source = (ROOT / "resource_miner" / "job_application.py").read_text()
        for name in ("urllib", "requests", "socket", "http.client", "smtplib"):
            self.assertNotIn(f"import {name}", source)


class SecondPassRegressionTests(unittest.TestCase):
    """Added after the independent diff review of ae772bd8."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_dot_git_directory_counts_as_inside_a_repository(self):
        from resource_miner.application_portfolio import _inside_git_work_tree
        repo = self.root / "other"
        make_repo(repo, "https://github.com/humanaios-ui/opportunity-aggregator", {"a": "x"})
        self.assertTrue(_inside_git_work_tree(repo))
        self.assertTrue(_inside_git_work_tree(repo / ".git"))

    def test_plain_directory_outside_git_is_outside(self):
        from resource_miner.application_portfolio import _inside_git_work_tree
        probe = subprocess.run(["git", "-C", str(self.root), "rev-parse", "--is-inside-work-tree"],
                               capture_output=True, text=True)
        if probe.returncode == 0:
            self.skipTest("temp dir is inside a git work tree on this machine")
        self.assertFalse(_inside_git_work_tree(self.root))

    def test_mixed_in_progress_and_documented_cannot_be_supported(self):
        rec = record(evidence=[{"id": "e1", "claim": "c", "source": "s", "state": "DOCUMENTED"},
                               {"id": "e2", "claim": "c2", "source": "s", "state": "IN_PROGRESS"}],
                     requirements=[{"id": "r1", "requirement": "x", "kind": "duty", "status": "SUPPORTED",
                                    "mapping": "DIRECT", "evidence_ids": ["e1", "e2"]}])
        with self.assertRaises(ValueError):
            prepare(rec)

    def test_evidence_state_is_visible_in_the_map(self):
        html = prepare(record())["evidence_map_html"]
        self.assertIn("DOCUMENTED; source: doc-1", html)
        self.assertIn("USER_ATTESTED; source: doc-2", html)

    def test_contradicted_requirement_must_stay_a_gap_and_survives_source_removal(self):
        bad = record(requirements=[{"id": "r1", "requirement": "x", "kind": "duty",
                                    "status": "CONTRADICTED", "mapping": "DIRECT", "evidence_ids": ["e1"]}])
        with self.assertRaises(ValueError):
            prepare(bad)
        ok = record(requirements=[{"id": "r1", "requirement": "x", "kind": "duty",
                                   "status": "CONTRADICTED", "mapping": "GAP", "evidence_ids": ["e1"]}])
        out = downgrade_for_removed_evidence(ok, ["e1"])
        self.assertEqual(out["requirements"][0]["status"], "CONTRADICTED")

    def _consumer(self, body):
        repo = self.root / "consumer"
        make_repo(repo, "https://github.com/humanaios-ui/opportunity-aggregator",
                  {"lib/matching/humanaios-profile.ts": body})
        return repo

    GOOD = ("export const humanaiosProfile = {\n description: 'Open research.',\n"
            " fundingStage: 'research', budgetRange: null,\n teamProfile: { size: 'solo' },\n}\n"
            "/**\n * Calculate grant relevance\n */\n")

    def test_consumer_parse_accepts_the_expected_shape(self):
        data, _ = project_consumer(self._consumer(self.GOOD))
        self.assertEqual(data["description"], "Open research.")

    def test_commented_out_description_is_inert(self):
        body = self.GOOD.replace(" description:", " // description: 'OLD 374K verified workers'\n description:")
        data, _ = project_consumer(self._consumer(body))
        self.assertEqual(data["description"], "Open research.")

    def test_consumer_parse_fails_closed_on_ambiguity(self):
        for body in (
            self.GOOD.replace(" fundingStage:", " description: 'second',\n fundingStage:"),
            self.GOOD.replace(" fundingStage:", " /* description: 'x' */\n fundingStage:"),
            self.GOOD.replace("/**\n * Calculate grant relevance\n */\n", ""),
        ):
            with self.subTest(body=body[:40]):
                shutil.rmtree(self.root / "consumer", ignore_errors=True)
                with self.assertRaises(ValueError):
                    project_consumer(self._consumer(body))

    def test_profile_presence_cli_has_no_traceback_on_missing_checkout(self):
        result = subprocess.run(
            [sys.executable, "-m", "resource_miner.cli", "profile-presence",
             "--private-checkout", str(self.root / "absent"), "--owner-approved"],
            cwd=ROOT, capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertNotIn("Traceback", result.stderr)
        self.assertIn("repository checkout not found", result.stderr)


class ProfilePresenceRegressionTests(unittest.TestCase):
    def test_repeated_candidate_retains_every_source_id(self):
        from resource_miner.profile_presence import reconcile_public_candidates
        url = "https://github.com/example-handle"
        obs = [{"provider": "github", "url": url, "source_id": "src-one"},
               {"provider": "github", "url": url, "source_id": "src-two"}]
        out = reconcile_public_candidates([], obs, owner_approved=True)
        cands = out["candidates"]
        self.assertEqual(len(cands), 1)
        self.assertEqual(cands[0]["source_ids"], ["src-one", "src-two"])
        self.assertEqual(cands[0]["source_id"], "src-one")

    def test_orcid_null_identifier_is_inconclusive_not_a_crash(self):
        from resource_miner.profile_presence import public_readback
        seeds = [{"provider": "orcid", "url": "https://orcid.org/0000-0002-1825-0097"}]
        body = json.dumps({"orcid-identifier": None}).encode()
        result = public_readback(seeds, owner_approved=True, transport=lambda u: (200, body))
        self.assertEqual(result[0]["state"], "READBACK_INCONCLUSIVE")


if __name__ == "__main__":
    unittest.main()
