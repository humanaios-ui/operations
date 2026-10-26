"""Owner-only application document from a private master profile.

Source-of-truth: committed master_profile.yaml in humanaios-internal.
Format: the provided project/program resume's seven-section layout.
Never invent work history or persist profile data inside a repository.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

from .private_profile_bridge import load_private_master

DELIVERY = [
    ("DEFINE", "objective, constraints, acceptance criteria"),
    ("SCOPE", "work packages, stakeholders, dependencies"),
    ("RESOURCE", "people, budget, systems, evidence"),
    ("EXECUTE", "coordinate, implement, escalate"),
    ("VERIFY", "validate, audit, measure, reconcile"),
    ("CORRECT / HANDOFF", "close gaps, transfer ownership, retain evidence"),
]

ORDER = [
    "PROFESSIONAL PROFILE",
    "PROJECT DELIVERY MODEL",
    "CORE CAPABILITIES",
    "SELECTED PROJECT EVIDENCE PORTFOLIO",
    "PROFESSIONAL EXPERIENCE",
    "TECHNICAL / QUALITY TOOLKIT",
    "EDUCATION & SELECTED RESEARCH",
]


def _safe_text(value: Any, max_length: int = 500) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("missing source-backed text")
    value = " ".join(value.split())
    if len(value) > max_length:
        raise ValueError("source text exceeds limit")
    return value


def _keyword_score(text: str, target: str) -> int:
    words = {w for w in re.findall(r"[a-z0-9]{3,}", target.lower()) if w not in {
        "the", "and", "for", "with", "from", "role", "job", "position",
    }}
    lower = text.lower()
    return sum(1 for w in words if w in lower)


def _sort_records(records: list[tuple[str, list[str], str]], target: str) -> list:
    # Stable base order. Target can change order, never facts or provenance.
    return sorted(records, key=lambda row: -_keyword_score(" ".join([row[0], *row[1]]), target))


def assemble_application(master: dict[str, Any], *,
                         target: str = "", max_projects: int = 5) -> dict[str, Any]:
    """Pure content projection, no writes, no new achievement claims."""
    if not isinstance(master, dict) or not isinstance(master.get("contact"), dict):
        raise ValueError("invalid private master")
    if not isinstance(master.get("experience"), list) or not isinstance(master.get("achievements"), list):
        raise ValueError("missing source-backed history")
    contact = master["contact"]
    summary = master.get("summary") or {}
    professional = _safe_text(summary.get("professional_brief", ""), 1500)
    competencies = master.get("competencies") or {}
    capabilities = []
    for domain in ("operations", "data_quality", "ai_research"):
        values = competencies.get(domain) or []
        if not isinstance(values, list):
            raise ValueError("bad competency section")
        for item in values:
            capabilities.append(_safe_text(item))
    technical = competencies.get("technical_skills") or {}
    toolkit = []
    if isinstance(technical, dict):
        for domain, values in technical.items():
            if not isinstance(values, list):
                continue
            toolkit.append(f"{domain.replace('_', ' ').title()}: " + ", ".join(
                _safe_text(x, 100) for x in values[:16]
            ))
    if not capabilities:
        raise ValueError("master has no competencies")

    projects = []
    for i, category in enumerate(master["achievements"]):
        if not isinstance(category, dict):
            raise ValueError("bad achievement category")
        for j, row in enumerate(category.get("items") or []):
            if not isinstance(row, dict):
                raise ValueError("bad evidence item")
            label = _safe_text(row.get("title"), 220)
            desc = _safe_text(row.get("description"), 500)
            projects.append((label, [desc], f"achievements[{i}].items[{j}]"))

    experience = []
    for i, row in enumerate(master["experience"]):
        if not isinstance(row, dict):
            raise ValueError("bad experience item")
        label = " | ".join([
            _safe_text(row.get("role"), 180),
            _safe_text(row.get("organization"), 180),
            str(row.get("start_date") or "")[:10] + " - " +
            str(row.get("end_date") or "Present")[:10],
        ])
        bullets = [_safe_text(x, 450) for x in (row.get("bullets") or [])][:3]
        experience.append({
            "heading": label,
            "description": _safe_text(row.get("summary"), 500),
            "bullets": bullets,
            "source": f"experience[{i}]",
        })

    edu = []
    for i, row in enumerate(master.get("education") or []):
        degree = _safe_text(row.get("degree"))
        school = _safe_text(row.get("school"))
        status = str(row.get("status") or row.get("year_completed") or "")
        edu.append({"text": f"{degree} - {school}" + (f" ({status})" if status else ""),
                    "source": f"education[{i}]"})

    research = []
    for i, row in enumerate(master.get("publications") or []):
        if len(research) >= 2:
            break
        text = _safe_text(row.get("title"), 220)
        status = str(row.get("status") or row.get("type") or "")
        research.append({"text": text + (f" ({status})" if status else ""),
                         "source": f"publications[{i}]"})

    for i, row in enumerate(master.get("project_evidence") or []):
        if not isinstance(row, dict):
            raise ValueError("bad project evidence")
        if row.get("review_status") != "APPROVED":
            continue
        if not row.get("evidence_ref"):
            raise ValueError("project evidence needs a source reference")
        label = _safe_text(row.get("title"), 220)
        mandate = _safe_text(row.get("mandate"), 500)
        bullets = row.get("evidence_bullets") or []
        if not isinstance(bullets, list) or not bullets or len(bullets) > 8:
            raise ValueError("invalid project evidence bullets")
        projects.append((label, [mandate] + [_safe_text(x, 450) for x in bullets],
                         f"project_evidence[{i}]"))

    selected = _sort_records(projects, target)[:max_projects]
    return {
        "schema": "humanaios.application-portfolio.v0.1",
        "format_reference": "HumanAIOS_Master_Project_Management_Resume.docx",
        "status": "DRAFT_REVIEW_REQUIRED",
        "authority_effect": "NONE",
        "target": target[:250],
        "source_id": _safe_text(master.get("profile_id"), 100),
        "sections": ORDER,
        "person": {
            "name": _safe_text(contact.get("legal_name"), 200),
            "location": _safe_text(contact.get("location"), 120),
            "email": _safe_text(contact.get("email"), 160),
            "phone": _safe_text(contact.get("phone"), 80),
        },
        "professional_profile": professional,
        "delivery_model": DELIVERY,
        "capabilities": _sort_records([(x, [], "") for x in capabilities], target),
        "projects": [
            {"heading": label, "bullets": bullets, "source": ref}
            for label, bullets, ref in selected
        ],
        "experience": experience,
        "toolkit": toolkit,
        "education_and_research": edu + research,
        "warnings": [
            "Template project evidence is not yet incorporated into the private master.",
            "Only master-backed facts appear; no target-specific claims are invented.",
            "Review all dates and publication statuses before submission.",
        ],
    }


def render_docx(application: dict[str, Any], template: Path, output: Path,
                *, protected_roots: tuple[Path, ...] = ()) -> None:
    """Render privately, borrowing the user's DOCX page and heading styles."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    template = Path(template).expanduser().resolve(strict=True)
    output = Path(output).expanduser().resolve()
    for root in protected_roots:
        if output == root.resolve() or root.resolve() in output.parents:
            raise ValueError("refusing to write a private application into a git checkout")
    if output.exists():
        raise FileExistsError("will not overwrite an existing application")
    if not output.parent.is_dir():
        raise ValueError("output parent must exist")
    doc = Document(str(template))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for child in list(container._element):
                if child.tag != qn("w:sectPr"):
                    container._element.remove(child)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""

    def normal(text: str, *, strong: bool = False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = strong
        return p

    def heading(text: str):
        p = doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)

    normal(application["person"]["name"].upper(), strong=True)
    normal("PROJECT / PROGRAM DELIVERY | QUALITY SYSTEMS | AI RESEARCH INFRASTRUCTURE | OPERATIONS")
    c = application["person"]
    normal(f'{c["location"]} | {c["email"]} | {c["phone"]}')
    heading("PROFESSIONAL PROFILE")
    normal(application["professional_profile"])
    heading("PROJECT DELIVERY MODEL")
    table = doc.add_table(rows=2, cols=6)
    for col, (label, detail) in enumerate(application["delivery_model"]):
        table.cell(0, col).text = label
        table.cell(1, col).text = detail
    heading("CORE CAPABILITIES")
    for label, _, _ in application["capabilities"][:8]:
        doc.add_paragraph(label, style="List Bullet")
    heading("SELECTED PROJECT EVIDENCE PORTFOLIO")
    for item in application["projects"]:
        normal(item["heading"], strong=True)
        for bullet in item["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")
    heading("PROFESSIONAL EXPERIENCE")
    for item in application["experience"]:
        normal(item["heading"], strong=True)
        normal(item["description"])
    heading("TECHNICAL / QUALITY TOOLKIT")
    for item in application["toolkit"][:5]:
        normal(item)
    heading("EDUCATION & SELECTED RESEARCH")
    for item in application["education_and_research"]:
        doc.add_paragraph(item["text"], style="List Bullet")
    old_umask = os.umask(0o077)
    try:
        doc.save(str(output))
        os.chmod(output, 0o600)
    finally:
        os.umask(old_umask)


def build_application(private_checkout: str | Path, template: str | Path,
                      output: str | Path, *, target: str = "", max_projects: int = 5,
                      forbidden_checkouts: tuple[Path, ...] = ()) -> dict[str, Any]:
    master, revision = load_private_master(private_checkout)
    app = assemble_application(master, target=target, max_projects=max_projects)
    roots = (Path(private_checkout).expanduser().resolve(), *forbidden_checkouts)
    render_docx(app, Path(template), Path(output), protected_roots=roots)
    return {
        "status": "PRIVATE_DRAFT_CREATED",
        "source_revision": revision,
        "selected_project_count": len(app["projects"]),
        "experience_count": len(app["experience"]),
        "warnings": app["warnings"],
        "authority_effect": "NONE",
    }
